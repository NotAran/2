"""Cajas por turno.

Hay varias cajas (Caja 1, Caja 2, ...). Cada día tiene 3 turnos (noche, mañana y
tarde) y, para cada caja y turno, una planilla de billeteras (filas, compartidas
por todas las cajas) por personas (columnas, propias de cada caja):

- Turno actual: la suma de los totales de las columnas de esa caja.
- Turno anterior: el turno actual de la misma caja en el turno anterior.

Resultado = turno anterior − (turno actual + bajada), final = resultado + saldo
y total final = final − bono. Cada caja tiene sus propios montos, panel y
observaciones, y una caja no ve ni altera los cálculos de otra. El admin define
cajas, billeteras y columnas, crea usuarios y les asigna qué cajas ven y editan.
Los datos se guardan en una base SQLite interna y se pueden descargar en Excel.
"""

import io
import os
import secrets
import sqlite3
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

from flask import Flask, abort, flash, g, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
VERSION_ESQUEMA = 9  # la 9 agregó colores a las billeteras; se quitaron, pero las bases que ya la tienen quedan igual
LARGO_OBSERVACIONES = 2000
CAJAS_INICIALES = ("Caja 1", "Caja 2")
BILLETERAS_INICIALES = ("Mercado", "Naranja", "Ualá", "Personal", "Brubank", "Lemon", "Prex", "Arq", "Binance")
COLUMNAS_INICIALES = (("Paco", 1), ("Antonio", 1), ("Ibra", 1))  # de la Caja 1; signo: suma o resta
TURNOS = (("noche", "Noche"), ("manana", "Mañana"), ("tarde", "Tarde"))  # orden de los turnos en el día
# Horario de cada turno. La noche de un día arranca a las 23:00 del día anterior,
# y el día termina (y sus turnos se cierran) a las 23:00, cuando termina la tarde.
HORARIOS = ("23:00 a 07:00", "07:00 a 15:00", "15:00 a 23:00")
HORA_CIERRE = time(23, 0)
ZONA_HORARIA = ZoneInfo(os.environ.get("MONTOS_TZ", "America/Argentina/Buenos_Aires"))
ROLES = (("admin", "Admin"), ("encargado", "Encargado"), ("cajero", "Cajero"))
TURNO_POR_SLUG = {slug: i for i, (slug, _) in enumerate(TURNOS)}
# Datos del panel de cada caja y turno, en el orden en que se muestran.
# El saldo no se carga: se calcula como depósito − retiro.
EXTRAS = (("deposito", "Depósito"), ("retiro", "Retiro"), ("bono", "Bono"), ("saldo", "Saldo"), ("bajada", "Bajada"))
EXTRAS_CALCULADOS = {"saldo"}
DIAS = ("lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo")
FORMATO_MONEDA = '"$"#,##0.00;[Red]-"$"#,##0.00'


def clave_secreta():
    """Usa SECRET_KEY si está definida; si no, genera una y la guarda junto a la app."""
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    ruta = os.path.join(BASE_DIR, ".secret_key")
    if not os.path.exists(ruta):
        with open(ruta, "w") as f:
            f.write(secrets.token_hex(32))
    with open(ruta) as f:
        return f.read().strip()


app = Flask(__name__)
app.config.from_mapping(
    SECRET_KEY=clave_secreta(),
    DATABASE=os.environ.get("MONTOS_DB", os.path.join(BASE_DIR, "montos.db")),
)


# --- Base de datos -----------------------------------------------------------
# La tabla "filas" guarda las billeteras.

ESQUEMA = """
CREATE TABLE IF NOT EXISTS usuarios (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    nombre     TEXT    NOT NULL UNIQUE COLLATE NOCASE,
    clave_hash TEXT    NOT NULL,
    rol        TEXT    NOT NULL DEFAULT 'cajero',  -- clave en ROLES
    nombre_visible TEXT  -- nombre que elige el usuario; si no hay, se usa el de ingreso
);
CREATE TABLE IF NOT EXISTS cajas (
    id     INTEGER PRIMARY KEY AUTOINCREMENT,
    nombre TEXT    NOT NULL UNIQUE COLLATE NOCASE,
    orden  INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS filas (
    id     INTEGER PRIMARY KEY AUTOINCREMENT,
    nombre TEXT    NOT NULL UNIQUE COLLATE NOCASE,
    orden  INTEGER NOT NULL DEFAULT 0,
    activo INTEGER NOT NULL DEFAULT 1  -- 0 = quitada pero con historial
);
CREATE TABLE IF NOT EXISTS columnas (
    id      INTEGER PRIMARY KEY AUTOINCREMENT,
    caja_id INTEGER NOT NULL REFERENCES cajas(id) ON DELETE CASCADE,
    nombre  TEXT    NOT NULL COLLATE NOCASE,
    signo   INTEGER NOT NULL DEFAULT 1 CHECK (signo IN (1, -1)),
    orden   INTEGER NOT NULL DEFAULT 0,
    activo  INTEGER NOT NULL DEFAULT 1,  -- 0 = quitada pero con historial
    UNIQUE (caja_id, nombre)
);
CREATE TABLE IF NOT EXISTS celdas (
    fecha           TEXT    NOT NULL,  -- AAAA-MM-DD
    turno           INTEGER NOT NULL CHECK (turno IN (0, 1, 2)),  -- índice en TURNOS
    fila_id         INTEGER NOT NULL REFERENCES filas(id)    ON DELETE CASCADE,
    columna_id      INTEGER NOT NULL REFERENCES columnas(id) ON DELETE CASCADE,  -- la columna define la caja
    centavos        INTEGER NOT NULL,  -- monto en centavos, sin decimales flotantes
    actualizado_por TEXT,
    actualizado     TEXT,
    PRIMARY KEY (fecha, turno, fila_id, columna_id)
);
CREATE TABLE IF NOT EXISTS extras (
    caja_id         INTEGER NOT NULL REFERENCES cajas(id) ON DELETE CASCADE,
    fecha           TEXT    NOT NULL,
    turno           INTEGER NOT NULL CHECK (turno IN (0, 1, 2)),
    campo           TEXT    NOT NULL,  -- clave en EXTRAS
    centavos        INTEGER NOT NULL,
    actualizado_por TEXT,
    actualizado     TEXT,
    PRIMARY KEY (caja_id, fecha, turno, campo)
);
CREATE TABLE IF NOT EXISTS observaciones (
    caja_id         INTEGER NOT NULL REFERENCES cajas(id) ON DELETE CASCADE,
    fecha           TEXT    NOT NULL,
    turno           INTEGER NOT NULL CHECK (turno IN (0, 1, 2)),
    texto           TEXT    NOT NULL,
    actualizado_por TEXT,
    actualizado     TEXT,
    PRIMARY KEY (caja_id, fecha, turno)
);
CREATE TABLE IF NOT EXISTS permisos_cajas (
    usuario_id INTEGER NOT NULL REFERENCES usuarios(id) ON DELETE CASCADE,
    caja_id    INTEGER NOT NULL REFERENCES cajas(id)    ON DELETE CASCADE,
    PRIMARY KEY (usuario_id, caja_id)
);
"""


def columnas_de(db, tabla):
    return {r["name"] for r in db.execute(f"PRAGMA table_info({tabla})")}


def reconstruir(db, tabla, crear, copiar):
    """Recrea una tabla con otra estructura (SQLite no permite cambiar claves con ALTER)."""
    db.executescript(f"""
        {crear.replace(f"CREATE TABLE IF NOT EXISTS {tabla} (", f"CREATE TABLE {tabla}_nueva (")}
        INSERT INTO {tabla}_nueva {copiar};
        DROP TABLE {tabla};
        ALTER TABLE {tabla}_nueva RENAME TO {tabla};
    """)


def definicion(tabla):
    ini = ESQUEMA.index(f"CREATE TABLE IF NOT EXISTS {tabla} (")
    return ESQUEMA[ini:ESQUEMA.index(");", ini) + 2]


def inicializar(db):
    """Crea o actualiza las tablas y carga los datos iniciales si no hay."""
    db.execute("PRAGMA foreign_keys = OFF")  # para poder recrear tablas sin borrar en cascada
    if db.execute("PRAGMA user_version").fetchone()[0] == 2:
        db.execute("DROP TABLE IF EXISTS celdas")  # versión anterior: montos sin caja ni turno
    db.executescript(ESQUEMA)
    columnas_usuarios = columnas_de(db, "usuarios")
    if "rol" not in columnas_usuarios:  # versiones anteriores: es_admin 1/0
        db.execute("ALTER TABLE usuarios ADD COLUMN rol TEXT NOT NULL DEFAULT 'cajero'")
        db.execute("UPDATE usuarios SET rol = 'admin' WHERE es_admin = 1")
    if "nombre_visible" not in columnas_usuarios:
        db.execute("ALTER TABLE usuarios ADD COLUMN nombre_visible TEXT")
    for tabla in ("filas", "columnas"):
        if "activo" not in columnas_de(db, tabla):
            db.execute(f"ALTER TABLE {tabla} ADD COLUMN activo INTEGER NOT NULL DEFAULT 1")
    if not db.execute("SELECT 1 FROM cajas").fetchone():
        db.executemany("INSERT INTO cajas (nombre, orden) VALUES (?, ?)", [(n, i) for i, n in enumerate(CAJAS_INICIALES)])
    caja_1 = db.execute("SELECT id FROM cajas ORDER BY orden, id LIMIT 1").fetchone()[0]

    # Versión 8: columnas, panel y observaciones pasan a ser de cada caja; lo anterior queda en la primera.
    if "caja_id" not in columnas_de(db, "columnas"):
        reconstruir(db, "columnas", definicion("columnas"),
                    f"(id, caja_id, nombre, signo, orden, activo) SELECT id, {caja_1}, nombre, signo, orden, activo FROM columnas")
    for tabla, campos in (("extras", "fecha, turno, campo, centavos, actualizado_por, actualizado"),
                          ("observaciones", "fecha, turno, texto, actualizado_por, actualizado")):
        if "caja_id" not in columnas_de(db, tabla):
            reconstruir(db, tabla, definicion(tabla), f"(caja_id, {campos}) SELECT {caja_1}, {campos} FROM {tabla}")
    tablas = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
    if "permisos_columnas" in tablas:  # antes se asignaban columnas: esos usuarios pasan a tener la primera caja
        db.execute(f"INSERT OR IGNORE INTO permisos_cajas SELECT DISTINCT usuario_id, {caja_1} FROM permisos_columnas")
    db.execute("DROP TABLE IF EXISTS permisos_columnas")
    db.execute("DROP TABLE IF EXISTS permisos_filas")

    db.execute("DELETE FROM extras WHERE campo = 'saldo'")  # desde la versión 5 el saldo se calcula
    if not db.execute("SELECT 1 FROM filas").fetchone():
        db.executemany(
            "INSERT INTO filas (nombre, orden) VALUES (?, ?)", [(n, i) for i, n in enumerate(BILLETERAS_INICIALES)]
        )
    if not db.execute("SELECT 1 FROM columnas").fetchone():
        db.executemany(
            "INSERT INTO columnas (caja_id, nombre, signo, orden) VALUES (?, ?, ?, ?)",
            [(caja_1, n, signo, i) for i, (n, signo) in enumerate(COLUMNAS_INICIALES)],
        )
    db.execute(f"PRAGMA user_version = {VERSION_ESQUEMA}")
    db.commit()
    db.execute("PRAGMA foreign_keys = ON")


def get_db():
    if "db" not in g:
        db = sqlite3.connect(app.config["DATABASE"])
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys = ON")
        if db.execute("PRAGMA user_version").fetchone()[0] < VERSION_ESQUEMA:
            inicializar(db)
        g.db = db
    return g.db


@app.teardown_appcontext
def cerrar_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


# --- Utilidades --------------------------------------------------------------

def parsear_monto(texto):
    """Convierte texto como '1.234,56', '1234.56', '-500' o '$ 500' a centavos (int).

    Si aparecen punto y coma, el último que aparece es el separador decimal.
    Si aparece uno solo, es de miles cuando se repite o lo siguen exactamente
    3 dígitos ('9.850' = 9850, '1,000' = 1000); si no, es decimal ('12,5').
    """
    t = (texto or "").strip().replace("$", "").replace(" ", "")
    negativo = t.startswith("-")
    t = t.removeprefix("-")
    if "," in t and "." in t:
        decimal_sep = "," if t.rfind(",") > t.rfind(".") else "."
        miles_sep = "." if decimal_sep == "," else ","
        t = t.replace(miles_sep, "").replace(decimal_sep, ".")
    elif "," in t or "." in t:
        sep = "," if "," in t else "."
        entero, _, dec = t.rpartition(sep)
        if t.count(sep) > 1 or len(dec) == 3:
            t = t.replace(sep, "")
        else:
            t = f"{entero}.{dec}"
    try:
        valor = Decimal(t)
    except InvalidOperation:
        raise ValueError("no es un número válido")
    if not valor.is_finite() or valor < 0:
        raise ValueError("no es un número válido")
    centavos = (valor * 100).quantize(Decimal("1"))
    if centavos != valor * 100:
        raise ValueError("admite como máximo 2 decimales")
    return -int(centavos) if negativo else int(centavos)


@app.template_filter("numero")
def formato_numero(centavos):
    """123456 -> '1.234,56' (formato argentino, sin símbolo)."""
    signo = "-" if centavos < 0 else ""
    entero, resto = divmod(abs(int(centavos)), 100)
    return f"{signo}{entero:,}".replace(",", ".") + f",{resto:02d}"


@app.template_filter("dinero")
def formato_dinero(centavos):
    numero = formato_numero(centavos)
    return f"-${numero[1:]}" if numero.startswith("-") else f"${numero}"


def validar_usuario(nombre, clave):
    if not 3 <= len(nombre) <= 30 or not nombre.replace("_", "").replace(".", "").isalnum():
        return "El usuario debe tener entre 3 y 30 letras o números (sin espacios)."
    return validar_clave(clave)


def validar_clave(clave):
    if len(clave) < 6:
        return "La contraseña debe tener al menos 6 caracteres."
    return None


@app.template_global()
def es_admin(usuario):
    return usuario is not None and usuario["rol"] == "admin"


@app.template_global()
def nombre_de(usuario):
    """Nombre que se muestra: el que eligió el usuario, o su usuario de ingreso."""
    return usuario["nombre_visible"] or usuario["nombre"]


@app.template_global()
def nombre_rol(rol):
    return dict(ROLES).get(rol, rol)


def limpiar_nombre_visible(texto):
    return " ".join((texto or "").split())[:40] or None


def ahora():
    """Fecha y hora actual en la zona horaria del negocio."""
    return datetime.now(ZONA_HORARIA)


# --- Sesión, permisos y protección CSRF ----------------------------------------

def usuario_actual():
    if "usuario" not in g:
        uid = session.get("usuario_id")
        g.usuario = get_db().execute("SELECT * FROM usuarios WHERE id = ?", (uid,)).fetchone() if uid else None
    return g.usuario


def hay_usuarios():
    return get_db().execute("SELECT 1 FROM usuarios LIMIT 1").fetchone() is not None


@app.before_request
def proteger():
    endpoint = request.endpoint or ""
    if endpoint == "static":
        return None
    token = session.setdefault("csrf", secrets.token_hex(16))
    if request.method == "POST":
        if not secrets.compare_digest(token, request.form.get("csrf", "")):
            abort(400, "El formulario venció. Recargá la página e intentá de nuevo.")
    if not hay_usuarios():
        return None if endpoint == "setup" else redirect(url_for("setup"))
    if endpoint in ("login", "setup"):
        return None
    if usuario_actual() is None:
        return redirect(url_for("login"))
    if endpoint.startswith("admin") and not es_admin(usuario_actual()):
        abort(403)
    return None


@app.context_processor
def contexto():
    return {"csrf": session.get("csrf", ""), "usuario": usuario_actual()}


# --- Turnos y cajas ----------------------------------------------------------
# Un turno de un día ("periodo") se identifica por (fecha, turno), con turno = índice en TURNOS.
# Una caja es una fila de la tabla cajas.

def mover(periodo, pasos):
    """Periodo que está `pasos` turnos después (o antes, si es negativo)."""
    fecha, turno = periodo
    n = turno + pasos
    return fecha + timedelta(days=n // len(TURNOS)), n % len(TURNOS)


@app.template_global()
def nombre_periodo(periodo):
    fecha, turno = periodo
    return f"{DIAS[fecha.weekday()].capitalize()} {fecha:%d/%m/%Y} – {TURNOS[turno][1]}"


@app.template_global()
def url_caja(caja_id, periodo, endpoint="planilla"):
    fecha, turno = periodo
    return url_for(endpoint, caja=caja_id, fecha=fecha.isoformat(), turno=TURNOS[turno][0])


@app.template_global()
def turno_slug(periodo):
    return TURNOS[periodo[1]][0]


def periodo_en_curso(momento=None):
    """Turno que está corriendo: noche 23–7, mañana 7–15, tarde 15–23."""
    momento = momento or ahora()
    hora, hoy = momento.hour, momento.date()
    if hora >= 23:
        return hoy + timedelta(days=1), 0  # la noche ya pertenece al día siguiente
    if hora < 7:
        return hoy, 0
    return (hoy, 1) if hora < 15 else (hoy, 2)


def dia_cerrado(periodo, momento=None):
    """El día de un turno termina a las 23:00 de su fecha; después solo el admin puede cambiarlo."""
    momento = momento or ahora()
    return momento >= datetime.combine(periodo[0], HORA_CIERRE, tzinfo=momento.tzinfo or ZONA_HORARIA)


def periodo_pedido(estricto=False):
    """Turno de los parámetros fecha y turno; si faltan, el turno en curso."""
    try:
        return date.fromisoformat(request.values.get("fecha", "")), TURNO_POR_SLUG[request.values.get("turno", "")]
    except (ValueError, KeyError):
        if estricto:
            abort(400, "Turno inválido.")
        return periodo_en_curso()


def cajas_visibles(db, usuario):
    """Cajas que el usuario puede ver y editar: todas si es admin, si no las asignadas."""
    if es_admin(usuario):
        return db.execute("SELECT * FROM cajas ORDER BY orden, id").fetchall()
    return db.execute(
        """SELECT c.* FROM cajas c JOIN permisos_cajas p ON p.caja_id = c.id
           WHERE p.usuario_id = ? ORDER BY c.orden, c.id""",
        (usuario["id"],),
    ).fetchall()


def caja_pedida(db, usuario):
    """Caja del parámetro caja (tiene que ser visible para el usuario); si falta, la primera visible."""
    visibles = cajas_visibles(db, usuario)
    pedida = request.values.get("caja")
    if pedida is None:
        return (visibles[0] if visibles else None), visibles
    for caja in visibles:
        if str(caja["id"]) == pedida:
            return caja, visibles
    abort(403)


def turno_actual(db, caja_id, periodo):
    """Suma (con signo) de todos los montos de una caja en un turno."""
    return db.execute(
        """SELECT COALESCE(SUM(ce.centavos * co.signo), 0) FROM celdas ce
           JOIN columnas co ON co.id = ce.columna_id
           WHERE co.caja_id = ? AND ce.fecha = ? AND ce.turno = ?""",
        (caja_id, periodo[0].isoformat(), periodo[1]),
    ).fetchone()[0]


def calcular_totales(anterior, actual, montos_extras):
    """Cuenta de la caja: resultado = anterior − (actual + bajada); final = resultado + saldo;
    total final = final − bono."""
    bajada = montos_extras.get("bajada", 0)
    saldo = montos_extras.get("deposito", 0) - montos_extras.get("retiro", 0)
    bono = montos_extras.get("bono", 0)
    resultado = anterior - (actual + bajada)
    return {
        "anterior": anterior,
        "actual": actual,
        "bajada": bajada,
        "resultado": resultado,
        "saldo": saldo,
        "final": resultado + saldo,
        "bono": bono,
        "total_final": resultado + saldo - bono,
    }


def armar_caja(db, caja, periodo):
    """Datos y cálculos de una caja en un turno.

    - Turno actual: suma de los totales de las columnas de la caja.
    - Turno anterior: turno actual de la misma caja en el turno inmediatamente anterior.
    - Resultado = turno anterior − (turno actual + bajada); final = resultado + saldo
      (saldo = depósito − retiro); total final = final − bono.
    Se muestran las billeteras y columnas activas, y las quitadas solo si tienen
    montos en este turno.
    """
    fecha, turno = periodo[0].isoformat(), periodo[1]
    celdas = {
        (r["fila_id"], r["columna_id"]): r
        for r in db.execute(
            """SELECT ce.* FROM celdas ce JOIN columnas co ON co.id = ce.columna_id
               WHERE co.caja_id = ? AND ce.fecha = ? AND ce.turno = ?""",
            (caja["id"], fecha, turno),
        )
    }
    filas = [
        f for f in db.execute("SELECT * FROM filas ORDER BY orden, id")
        if f["activo"] or any(k[0] == f["id"] for k in celdas)
    ]
    columnas = [
        c for c in db.execute("SELECT * FROM columnas WHERE caja_id = ? ORDER BY orden, id", (caja["id"],))
        if c["activo"] or any(k[1] == c["id"] for k in celdas)
    ]
    extras = {
        r["campo"]: r
        for r in db.execute(
            "SELECT * FROM extras WHERE caja_id = ? AND fecha = ? AND turno = ?", (caja["id"], fecha, turno)
        )
    }

    def valor(f, c):
        celda = celdas.get((f["id"], c["id"]))
        return celda["centavos"] if celda else 0

    # Se suma cada columna, y el turno actual es la suma de los totales de las columnas
    total_columna = {c["id"]: sum(valor(f, c) for f in filas) for c in columnas}
    return {
        "caja": caja,
        "periodo": periodo,
        "filas": filas,
        "columnas": columnas,
        "celdas": celdas,
        "total_columna": total_columna,
        "totales": calcular_totales(
            turno_actual(db, caja["id"], mover(periodo, -1)),
            turno_actual(db, caja["id"], periodo),
            {campo: r["centavos"] for campo, r in extras.items()},
        ),
        "extras": extras,
        "observaciones": db.execute(
            "SELECT * FROM observaciones WHERE caja_id = ? AND fecha = ? AND turno = ?", (caja["id"], fecha, turno)
        ).fetchone(),
    }


def historial(db, caja_id):
    """Una entrada por turno con datos en la caja, con sus totales (más nuevos primero)."""
    actuales = {
        (date.fromisoformat(r["fecha"]), r["turno"]): r["actual"]
        for r in db.execute(
            """SELECT ce.fecha, ce.turno, SUM(ce.centavos * co.signo) AS actual
               FROM celdas ce JOIN columnas co ON co.id = ce.columna_id
               WHERE co.caja_id = ? GROUP BY ce.fecha, ce.turno""",
            (caja_id,),
        )
    }
    extras = {}
    for r in db.execute("SELECT fecha, turno, campo, centavos FROM extras WHERE caja_id = ?", (caja_id,)):
        extras.setdefault((date.fromisoformat(r["fecha"]), r["turno"]), {})[r["campo"]] = r["centavos"]
    return [
        {"periodo": p, **calcular_totales(actuales.get(mover(p, -1), 0), actuales.get(p, 0), extras.get(p, {}))}
        for p in sorted(set(actuales) | set(extras), reverse=True)
    ]


@app.get("/")
def planilla():
    db = get_db()
    usuario = usuario_actual()
    caja, cajas = caja_pedida(db, usuario)
    periodo = periodo_pedido()
    contexto_turno = {
        "periodo": periodo,
        "periodo_previo": mover(periodo, -1),
        "periodo_siguiente": mover(periodo, 1),
        "turnos": [((periodo[0], i), nombre) for i, (_, nombre) in enumerate(TURNOS)],
        "horarios": HORARIOS,
        "en_curso": periodo_en_curso(),
        "cerrada": dia_cerrado(periodo),
        "bloqueada": dia_cerrado(periodo) and not es_admin(usuario),
    }
    if caja is None:
        return render_template("planilla.html", caja=None, cajas=[], **contexto_turno)
    return render_template(
        "planilla.html",
        **{**armar_caja(db, caja, periodo), **contexto_turno},
        cajas=cajas,
        historial=historial(db, caja["id"])[:15],
        campos_extras=EXTRAS,
        largo_observaciones=LARGO_OBSERVACIONES,
    )


def leer_cambio(texto, actual, etiqueta, errores):
    """Interpreta lo escrito en un campo. Devuelve (hay_cambio, centavos); centavos None = borrar."""
    texto = texto.strip()
    if not texto:
        return actual is not None, None
    try:
        centavos = parsear_monto(texto)
    except ValueError as e:
        errores.append(f"{etiqueta}: «{texto}» {e}.")
        return False, None
    return not (actual and actual["centavos"] == centavos), centavos


@app.post("/guardar")
def guardar():
    db = get_db()
    usuario = usuario_actual()
    if "caja" not in request.form:
        abort(400, "Falta la caja.")
    caja, _ = caja_pedida(db, usuario)  # 403 si el usuario no tiene esa caja
    periodo = periodo_pedido(estricto=True)
    if dia_cerrado(periodo) and not es_admin(usuario):
        flash(f"El día {periodo[0]:%d/%m/%Y} ya terminó: solo el admin puede hacer cambios.", "error")
        return redirect(url_caja(caja["id"], periodo))
    datos = armar_caja(db, caja, periodo)  # solo se guardan celdas de las columnas de esta caja
    fecha, turno = periodo[0].isoformat(), periodo[1]
    momento = ahora().strftime("%Y-%m-%d %H:%M")
    errores, cambios = [], 0

    for f in datos["filas"]:
        for c in datos["columnas"]:
            campo = f"c_{f['id']}_{c['id']}"
            if campo not in request.form:
                continue
            clave = (fecha, turno, f["id"], c["id"])
            hay_cambio, centavos = leer_cambio(
                request.form[campo], datos["celdas"].get((f["id"], c["id"])), f"{f['nombre']} / {c['nombre']}", errores
            )
            if not hay_cambio:
                continue
            if centavos is None:
                db.execute("DELETE FROM celdas WHERE fecha = ? AND turno = ? AND fila_id = ? AND columna_id = ?", clave)
            else:
                db.execute(
                    """INSERT INTO celdas (fecha, turno, fila_id, columna_id, centavos, actualizado_por, actualizado)
                       VALUES (?, ?, ?, ?, ?, ?, ?)
                       ON CONFLICT (fecha, turno, fila_id, columna_id) DO UPDATE SET
                         centavos = excluded.centavos,
                         actualizado_por = excluded.actualizado_por,
                         actualizado = excluded.actualizado""",
                    (*clave, centavos, nombre_de(usuario), momento),
                )
            cambios += 1

    clave_turno = (caja["id"], fecha, turno)
    for campo, nombre in EXTRAS:
        if campo in EXTRAS_CALCULADOS or f"x_{campo}" not in request.form:
            continue
        hay_cambio, centavos = leer_cambio(request.form[f"x_{campo}"], datos["extras"].get(campo), nombre, errores)
        if not hay_cambio:
            continue
        if centavos is None:
            db.execute("DELETE FROM extras WHERE caja_id = ? AND fecha = ? AND turno = ? AND campo = ?", (*clave_turno, campo))
        else:
            db.execute(
                """INSERT INTO extras (caja_id, fecha, turno, campo, centavos, actualizado_por, actualizado)
                   VALUES (?, ?, ?, ?, ?, ?, ?)
                   ON CONFLICT (caja_id, fecha, turno, campo) DO UPDATE SET
                     centavos = excluded.centavos,
                     actualizado_por = excluded.actualizado_por,
                     actualizado = excluded.actualizado""",
                (*clave_turno, campo, centavos, nombre_de(usuario), momento),
            )
        cambios += 1

    if "observaciones" in request.form:
        texto = request.form["observaciones"].replace("\r\n", "\n").strip()[:LARGO_OBSERVACIONES]
        anterior = datos["observaciones"]["texto"] if datos["observaciones"] else ""
        if texto != anterior:
            if texto:
                db.execute(
                    """INSERT INTO observaciones (caja_id, fecha, turno, texto, actualizado_por, actualizado)
                       VALUES (?, ?, ?, ?, ?, ?)
                       ON CONFLICT (caja_id, fecha, turno) DO UPDATE SET
                         texto = excluded.texto,
                         actualizado_por = excluded.actualizado_por,
                         actualizado = excluded.actualizado""",
                    (*clave_turno, texto, nombre_de(usuario), momento),
                )
            else:
                db.execute("DELETE FROM observaciones WHERE caja_id = ? AND fecha = ? AND turno = ?", clave_turno)
            cambios += 1

    db.commit()
    for error in errores:
        flash(error, "error")
    if cambios:
        flash(f"{caja['nombre']} · {nombre_periodo(periodo)}: se guardaron {cambios} cambio(s).", "ok")
    elif not errores:
        flash("No había cambios para guardar.", "ok")
    return redirect(url_caja(caja["id"], periodo))


def estilo_encabezado(ws, fila, cantidad):
    for i in range(1, cantidad + 1):
        celda = ws.cell(row=fila, column=i)
        celda.font, celda.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5597")


@app.get("/exportar.xlsx")
def exportar():
    """Excel con la caja y el turno elegidos (totales como fórmulas) y el historial de esa caja."""
    db = get_db()
    caja, _ = caja_pedida(db, usuario_actual())
    if caja is None:
        abort(404)
    periodo = periodo_pedido()
    datos = armar_caja(db, caja, periodo)
    filas, columnas, celdas, totales = datos["filas"], datos["columnas"], datos["celdas"], datos["totales"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Caja"
    ws["A1"] = f"{caja['nombre']}: {nombre_periodo(periodo)}"
    ws["A1"].font = Font(bold=True, size=13)
    encabezados = ["Billetera"] + [f"{c['nombre']}{'' if c['signo'] > 0 else ' (−)'}" for c in columnas]
    ws.append([])
    ws.append(encabezados)
    estilo_encabezado(ws, 3, len(encabezados))
    for i in range(1, len(encabezados) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16 if i > 1 else 18
    ws.freeze_panes = "B4"

    col_ultima = len(columnas) + 1
    letra = {c["id"]: get_column_letter(j) for j, c in enumerate(columnas, start=2)}
    for i, f in enumerate(filas, start=4):
        ws.cell(row=i, column=1, value=f["nombre"])
        for c in columnas:
            celda = celdas.get((f["id"], c["id"]))
            ws[f"{letra[c['id']]}{i}"] = celda["centavos"] / 100 if celda else None

    fila_total = len(filas) + 4
    ws.cell(row=fila_total, column=1, value="Total").font = Font(bold=True)
    for j in range(2, col_ultima + 1):
        l = get_column_letter(j)
        ws.cell(row=fila_total, column=j, value=f"=SUM({l}4:{l}{fila_total - 1})" if filas else 0).font = Font(
            bold=True
        )

    # Cuentas de la caja: el turno actual es la suma de los totales de las columnas
    r = fila_total + 2
    fila_extra = {campo: k for k, (campo, _) in enumerate(EXTRAS, start=r + 10)}
    terminos = "".join(f"{'+' if c['signo'] > 0 else '-'}{letra[c['id']]}{fila_total}" for c in columnas)
    resumen = (
        ("Turno anterior", totales["anterior"] / 100),
        ("Turno actual", f"={terminos or 0}"),
        ("Bajada", f"=B{fila_extra['bajada']}"),
        ("Resultado", f"=B{r}-(B{r + 1}+B{r + 2})"),
        ("Saldo", f"=B{fila_extra['saldo']}"),
        ("Final", f"=B{r + 3}+B{r + 4}"),
        ("Bono", f"=B{fila_extra['bono']}"),
        ("Total final", f"=B{r + 5}-B{r + 6}"),
    )
    for k, (nombre, valor) in enumerate(resumen, start=r):
        ws.cell(row=k, column=1, value=nombre)
        ws.cell(row=k, column=2, value=valor)
    for k in (r + 3, r + 5, r + 7):
        ws.cell(row=k, column=1).font = ws.cell(row=k, column=2).font = Font(bold=True)
    ws.cell(row=r + 9, column=1, value="Panel").font = Font(bold=True)
    for campo, nombre in EXTRAS:
        extra = datos["extras"].get(campo)
        ws.cell(row=fila_extra[campo], column=1, value=nombre)
        ws.cell(row=fila_extra[campo], column=2, value=extra["centavos"] / 100 if extra else None)
    ws.cell(row=fila_extra["saldo"], column=2, value=f"=B{fila_extra['deposito']}-B{fila_extra['retiro']}")
    fila_obs = fila_extra[EXTRAS[-1][0]] + 2
    ws.cell(row=fila_obs, column=1, value="Observaciones").font = Font(bold=True)
    if datos["observaciones"]:
        ws.cell(row=fila_obs + 1, column=1, value=datos["observaciones"]["texto"])
    for fila in ws.iter_rows(min_row=4, min_col=2, max_col=max(col_ultima, 2)):
        for c in fila:
            c.number_format = FORMATO_MONEDA

    hs = wb.create_sheet("Historial")
    campos = ("anterior", "actual", "bajada", "resultado", "saldo", "final", "bono", "total_final")
    hs.append(["Fecha", "Turno", "Turno anterior", "Turno actual", "Bajada", "Resultado", "Saldo", "Final", "Bono",
               "Total final"])
    estilo_encabezado(hs, 1, 10)
    for i in range(1, 11):
        hs.column_dimensions[get_column_letter(i)].width = 12 if i <= 2 else 16
    for h in reversed(historial(db, caja["id"])):
        hs.append([h["periodo"][0], TURNOS[h["periodo"][1]][1]] + [h[c] / 100 for c in campos])
        hs.cell(row=hs.max_row, column=1).number_format = "dd/mm/yyyy"
        for col in range(3, 11):
            hs.cell(row=hs.max_row, column=col).number_format = FORMATO_MONEDA

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nombre_archivo = "_".join(caja["nombre"].lower().split())
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"{nombre_archivo}_{periodo[0].isoformat()}_{TURNOS[periodo[1]][0]}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# --- Login ---------------------------------------------------------------------

@app.route("/setup", methods=["GET", "POST"])
def setup():
    """Primera vez: crear la cuenta del admin."""
    if hay_usuarios():
        return redirect(url_for("login"))
    if request.method == "POST":
        nombre, clave = request.form.get("nombre", "").strip(), request.form.get("clave", "")
        error = validar_usuario(nombre, clave)
        if not error:
            db = get_db()
            cur = db.execute(
                "INSERT INTO usuarios (nombre, clave_hash, rol, nombre_visible) VALUES (?, ?, 'admin', ?)",
                (nombre, generate_password_hash(clave), limpiar_nombre_visible(request.form.get("nombre_visible"))),
            )
            db.commit()
            session.clear()
            session["usuario_id"] = cur.lastrowid
            flash(f"Listo, {nombre}: sos el admin.", "ok")
            return redirect(url_for("planilla"))
        flash(error, "error")
    return render_template("setup.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        nombre, clave = request.form.get("nombre", "").strip(), request.form.get("clave", "")
        u = get_db().execute("SELECT * FROM usuarios WHERE nombre = ?", (nombre,)).fetchone()
        if u and check_password_hash(u["clave_hash"], clave):
            session.clear()
            session["usuario_id"] = u["id"]
            return redirect(url_for("planilla"))
        flash("Usuario o contraseña incorrectos.", "error")
    return render_template("login.html")


@app.post("/salir")
def salir():
    session.clear()
    return redirect(url_for("login"))


@app.route("/cuenta", methods=["GET", "POST"])
def cuenta():
    """Cada usuario elige su nombre y puede cambiar su contraseña."""
    usuario = usuario_actual()
    if request.method == "POST":
        db = get_db()
        nueva = request.form.get("clave_nueva", "")
        if nueva:
            if not check_password_hash(usuario["clave_hash"], request.form.get("clave_actual", "")):
                flash("La contraseña actual no es correcta.", "error")
                return redirect(url_for("cuenta"))
            error = validar_clave(nueva)
            if error:
                flash(error, "error")
                return redirect(url_for("cuenta"))
            db.execute("UPDATE usuarios SET clave_hash = ? WHERE id = ?", (generate_password_hash(nueva), usuario["id"]))
        db.execute(
            "UPDATE usuarios SET nombre_visible = ? WHERE id = ?",
            (limpiar_nombre_visible(request.form.get("nombre_visible")), usuario["id"]),
        )
        db.commit()
        flash("Tus datos se guardaron.", "ok")
        return redirect(url_for("cuenta"))
    return render_template("cuenta.html")


# --- Panel de admin -------------------------------------------------------------

@app.get("/admin")
def admin():
    db = get_db()
    cajas = db.execute("SELECT * FROM cajas ORDER BY orden, id").fetchall()
    permisos = {}
    for p in db.execute("SELECT * FROM permisos_cajas"):
        permisos.setdefault(p["usuario_id"], set()).add(p["caja_id"])
    columnas = {
        caja["id"]: {
            "activas": db.execute(
                "SELECT * FROM columnas WHERE caja_id = ? AND activo = 1 ORDER BY orden, id", (caja["id"],)
            ).fetchall(),
            "quitadas": db.execute(
                "SELECT * FROM columnas WHERE caja_id = ? AND activo = 0 ORDER BY orden, id", (caja["id"],)
            ).fetchall(),
        }
        for caja in cajas
    }
    return render_template(
        "admin.html",
        usuarios=db.execute(
            "SELECT * FROM usuarios ORDER BY CASE rol WHEN 'admin' THEN 0 WHEN 'encargado' THEN 1 ELSE 2 END, nombre"
        ).fetchall(),
        roles=ROLES,
        cajas=cajas,
        permisos=permisos,
        filas=db.execute("SELECT * FROM filas WHERE activo = 1 ORDER BY orden, id").fetchall(),
        filas_quitadas=db.execute("SELECT * FROM filas WHERE activo = 0 ORDER BY orden, id").fetchall(),
        columnas=columnas,
    )


def volver_admin(seccion):
    return redirect(url_for("admin") + f"#{seccion}")


@app.post("/admin/usuarios")
def admin_crear_usuario():
    nombre, clave = request.form.get("nombre", "").strip(), request.form.get("clave", "")
    rol = request.form.get("rol", "cajero")
    error = validar_usuario(nombre, clave) or (None if rol in dict(ROLES) else "Nivel inválido.")
    if not error:
        db = get_db()
        try:
            db.execute(
                "INSERT INTO usuarios (nombre, clave_hash, rol, nombre_visible) VALUES (?, ?, ?, ?)",
                (nombre, generate_password_hash(clave), rol, limpiar_nombre_visible(request.form.get("nombre_visible"))),
            )
            db.commit()
            extra = "" if rol == "admin" else " Ahora asignale las cajas que puede editar."
            flash(f"Usuario «{nombre}» ({nombre_rol(rol).lower()}) creado.{extra}", "ok")
        except sqlite3.IntegrityError:
            error = "Ya existe un usuario con ese nombre."
    if error:
        flash(error, "error")
    return volver_admin("usuarios")


@app.post("/admin/usuarios/<int:uid>/eliminar")
def admin_eliminar_usuario(uid):
    db = get_db()
    if uid == usuario_actual()["id"] or db.execute(
        "DELETE FROM usuarios WHERE id = ? AND rol != 'admin'", (uid,)
    ).rowcount == 0:
        abort(404)
    db.commit()
    flash("Usuario eliminado.", "ok")
    return volver_admin("usuarios")


@app.post("/admin/usuarios/<int:uid>/clave")
def admin_cambiar_clave(uid):
    clave = request.form.get("clave", "")
    error = validar_clave(clave)
    if error:
        flash(error, "error")
        return volver_admin("usuarios")
    db = get_db()
    if db.execute("UPDATE usuarios SET clave_hash = ? WHERE id = ?", (generate_password_hash(clave), uid)).rowcount == 0:
        abort(404)
    db.commit()
    flash("Contraseña actualizada.", "ok")
    return volver_admin("usuarios")


@app.post("/admin/usuarios/<int:uid>/datos")
def admin_datos_usuario(uid):
    """El admin cambia el nivel y el nombre visible de un usuario (su propio nivel no)."""
    db = get_db()
    u = db.execute("SELECT * FROM usuarios WHERE id = ?", (uid,)).fetchone()
    if not u:
        abort(404)
    rol = request.form.get("rol", u["rol"])
    if rol not in dict(ROLES) or (uid == usuario_actual()["id"] and rol != u["rol"]):
        flash("No podés cambiar tu propio nivel.", "error")
        return volver_admin("usuarios")
    db.execute(
        "UPDATE usuarios SET rol = ?, nombre_visible = ? WHERE id = ?",
        (rol, limpiar_nombre_visible(request.form.get("nombre_visible")), uid),
    )
    db.commit()
    flash(f"Datos de «{u['nombre']}» guardados.", "ok")
    return volver_admin("usuarios")


@app.post("/admin/usuarios/<int:uid>/cajas")
def admin_permisos(uid):
    """Qué cajas puede ver y editar un usuario (el admin ve y edita todas)."""
    db = get_db()
    if not db.execute("SELECT 1 FROM usuarios WHERE id = ? AND rol != 'admin'", (uid,)).fetchone():
        abort(404)
    db.execute("DELETE FROM permisos_cajas WHERE usuario_id = ?", (uid,))
    db.executemany(
        "INSERT INTO permisos_cajas (usuario_id, caja_id) SELECT ?, id FROM cajas WHERE id = ?",
        [(uid, caja_id) for caja_id in request.form.getlist("cajas", type=int)],
    )
    db.commit()
    flash("Cajas asignadas.", "ok")
    return volver_admin(f"permisos-{uid}")


@app.post("/admin/cajas")
def admin_crear_caja():
    nombre = " ".join(request.form.get("nombre", "").split())[:40]
    if not nombre:
        flash("Escribí un nombre para la caja.", "error")
        return volver_admin("cajas")
    db = get_db()
    try:
        orden = db.execute("SELECT COALESCE(MAX(orden), -1) + 1 FROM cajas").fetchone()[0]
        db.execute("INSERT INTO cajas (nombre, orden) VALUES (?, ?)", (nombre, orden))
        db.commit()
        flash(f"«{nombre}» creada. Agregale columnas y asignala a los usuarios.", "ok")
    except sqlite3.IntegrityError:
        flash(f"Ya existe una caja llamada «{nombre}».", "error")
    return volver_admin("cajas")


@app.post("/admin/cajas/<int:caja_id>/renombrar")
def admin_renombrar_caja(caja_id):
    nombre = " ".join(request.form.get("nombre", "").split())[:40]
    db = get_db()
    if not nombre:
        flash("Escribí un nombre para la caja.", "error")
        return volver_admin("cajas")
    try:
        if db.execute("UPDATE cajas SET nombre = ? WHERE id = ?", (nombre, caja_id)).rowcount == 0:
            abort(404)
        db.commit()
        flash(f"La caja ahora se llama «{nombre}».", "ok")
    except sqlite3.IntegrityError:
        flash(f"Ya existe una caja llamada «{nombre}».", "error")
    return volver_admin("cajas")


def ambito(tabla, item):
    """Filtro SQL para el grupo donde se ordena y se valida el nombre: las billeteras
    son todas un grupo; las columnas, uno por caja."""
    if tabla == "filas":
        return "", ()
    return " AND caja_id = ?", (item["caja_id"],)


@app.post("/admin/<any(filas, columnas):tabla>")
def admin_agregar(tabla):
    nombre = " ".join(request.form.get("nombre", "").split())[:40]
    seccion = "billeteras" if tabla == "filas" else "columnas"
    if not nombre:
        flash("Escribí un nombre.", "error")
        return volver_admin(seccion)
    db = get_db()
    if tabla == "columnas":
        caja = db.execute("SELECT * FROM cajas WHERE id = ?", (request.form.get("caja_id", type=int),)).fetchone()
        if not caja:
            abort(400, "Caja inválida.")
        seccion = f"columnas-{caja['id']}"
    filtro, params = ("", ()) if tabla == "filas" else (" AND caja_id = ?", (caja["id"],))
    orden = db.execute(f"SELECT COALESCE(MAX(orden), -1) + 1 FROM {tabla} WHERE 1{filtro}", params).fetchone()[0]
    existente = db.execute(f"SELECT * FROM {tabla} WHERE nombre = ?{filtro}", (nombre, *params)).fetchone()
    tipo = "billetera" if tabla == "filas" else "columna"
    if existente and existente["activo"]:
        flash(f"Ya existe una {tipo} llamada «{existente['nombre']}».", "error")
        return volver_admin(seccion)
    if existente:
        db.execute(f"UPDATE {tabla} SET activo = 1, orden = ? WHERE id = ?", (orden, existente["id"]))
        flash(f"«{existente['nombre']}» se volvió a agregar, con su historial.", "ok")
    elif tabla == "filas":
        db.execute("INSERT INTO filas (nombre, orden) VALUES (?, ?)", (nombre, orden))
        flash(f"Billetera «{nombre}» agregada al final. Usá las flechas para moverla.", "ok")
    else:
        signo = -1 if request.form.get("signo") == "-1" else 1
        db.execute(
            "INSERT INTO columnas (caja_id, nombre, signo, orden) VALUES (?, ?, ?, ?)", (caja["id"], nombre, signo, orden)
        )
        flash(f"Columna «{nombre}» agregada a {caja['nombre']}.", "ok")
    db.commit()
    return volver_admin(seccion)


def seccion_de(tabla, item):
    return "billeteras" if tabla == "filas" else f"columnas-{item['caja_id']}"


@app.post("/admin/<any(filas, columnas):tabla>/<int:item_id>/mover")
def admin_mover(tabla, item_id):
    """Sube o baja una billetera (o columna) un lugar entre las activas de su grupo."""
    db = get_db()
    item = db.execute(f"SELECT * FROM {tabla} WHERE id = ?", (item_id,)).fetchone()
    if not item:
        abort(404)
    filtro, params = ambito(tabla, item)
    ids = [r["id"] for r in db.execute(f"SELECT id FROM {tabla} WHERE activo = 1{filtro} ORDER BY orden, id", params)]
    if item_id in ids:
        i = ids.index(item_id)
        j = i - 1 if request.form.get("direccion") == "arriba" else i + 1
        if 0 <= j < len(ids):
            ids[i], ids[j] = ids[j], ids[i]
            db.executemany(f"UPDATE {tabla} SET orden = ? WHERE id = ?", [(k, x) for k, x in enumerate(ids)])
            db.commit()
    return volver_admin(seccion_de(tabla, item))


@app.post("/admin/<any(filas, columnas):tabla>/<int:item_id>/eliminar")
def admin_eliminar(tabla, item_id):
    """Quita una billetera o columna. Si ya tiene montos, se oculta para no cambiar los cálculos."""
    db = get_db()
    item = db.execute(f"SELECT * FROM {tabla} WHERE id = ?", (item_id,)).fetchone()
    if not item:
        abort(404)
    campo = "fila_id" if tabla == "filas" else "columna_id"
    if db.execute(f"SELECT 1 FROM celdas WHERE {campo} = ? LIMIT 1", (item_id,)).fetchone():
        db.execute(f"UPDATE {tabla} SET activo = 0 WHERE id = ?", (item_id,))
        flash(f"«{item['nombre']}» se quitó. Sus montos de turnos anteriores se conservan para no cambiar los cálculos.", "ok")
    else:
        db.execute(f"DELETE FROM {tabla} WHERE id = ?", (item_id,))
        flash(f"«{item['nombre']}» se eliminó.", "ok")
    db.commit()
    return volver_admin(seccion_de(tabla, item))


@app.post("/admin/<any(filas, columnas):tabla>/<int:item_id>/reactivar")
def admin_reactivar(tabla, item_id):
    db = get_db()
    item = db.execute(f"SELECT * FROM {tabla} WHERE id = ?", (item_id,)).fetchone()
    if not item:
        abort(404)
    filtro, params = ambito(tabla, item)
    orden = db.execute(f"SELECT COALESCE(MAX(orden), -1) + 1 FROM {tabla} WHERE 1{filtro}", params).fetchone()[0]
    db.execute(f"UPDATE {tabla} SET activo = 1, orden = ? WHERE id = ?", (orden, item_id))
    db.commit()
    flash(f"«{item['nombre']}» se volvió a agregar, con su historial.", "ok")
    return volver_admin(seccion_de(tabla, item))


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=os.environ.get("DEBUG") == "1")
