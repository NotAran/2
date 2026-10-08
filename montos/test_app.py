import io
import sqlite3
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

import app as montos


AHORA = datetime(2026, 10, 8, 5, 0, tzinfo=montos.ZONA_HORARIA)  # jueves 08/10, turno noche en curso


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(montos, "ahora", lambda: AHORA)
    montos.app.config.update(TESTING=True, DATABASE=str(tmp_path / "test.db"))
    with montos.app.test_client() as c:
        yield c


def post(client, url, **datos):
    """POST con el token CSRF de la sesión."""
    client.get("/login")  # asegura que la sesión tenga token
    with client.session_transaction() as s:
        datos["csrf"] = s["csrf"]
    return client.post(url, data=datos, follow_redirects=True)


def crear_admin(client):
    return post(client, "/setup", nombre="admin", clave="secreto1")


def ingresar(client, nombre, clave):
    post(client, "/salir")
    return post(client, "/login", nombre=nombre, clave=clave)


def cargar(client, fecha, turno, **celdas):
    """Guarda montos en una caja. celdas: {"Fila__Columna": "monto"}."""
    f, c = ids("filas"), ids("columnas")
    datos = {}
    for clave, monto in celdas.items():
        fila, columna = clave.split("__")
        datos[f"c_{f[fila]}_{c[columna]}"] = monto
    return post(client, "/guardar", fecha=fecha, turno=turno, **datos)


def caja(fecha, turno, usuario="admin"):
    with montos.app.app_context():
        db = montos.get_db()
        u = db.execute("SELECT * FROM usuarios WHERE nombre = ?", (usuario,)).fetchone()
        return montos.armar_caja(db, u, (date.fromisoformat(fecha), montos.TURNO_POR_SLUG[turno]))


def ids(tabla):
    with montos.app.app_context():
        return {r["nombre"]: r["id"] for r in montos.get_db().execute(f"SELECT id, nombre FROM {tabla} ORDER BY id")}


# --- Formato de montos ---------------------------------------------------------

@pytest.mark.parametrize("texto,centavos", [
    ("1234.56", 123456), ("1.234,56", 123456), ("1,234.56", 123456),
    ("$ 500", 50000), ("12,5", 1250), ("1,000", 100000), ("0.01", 1),
    ("9.850", 985000), ("1.234.567", 123456700), ("45.780,50", 4578050), ("3.5", 350),
    ("-500", -50000), ("-1.234,56", -123456), ("0", 0),
])
def test_parsear_monto(texto, centavos):
    assert montos.parsear_monto(texto) == centavos


@pytest.mark.parametrize("texto", ["", "abc", "--5", "1.2345", "nan", "inf"])
def test_parsear_monto_invalido(texto):
    with pytest.raises(ValueError):
        montos.parsear_monto(texto)


def test_formatos():
    assert montos.formato_dinero(123456789) == "$1.234.567,89"
    assert montos.formato_dinero(-50) == "-$0,50"
    # lo que se muestra en la celda se vuelve a leer igual
    for centavos in (0, 5, 123400, -98765432):
        assert montos.parsear_monto(montos.formato_numero(centavos)) == centavos


# --- Login y admin -----------------------------------------------------------

def test_primera_vez_pide_crear_admin(client):
    assert client.get("/").headers["Location"].endswith("/setup")
    r = crear_admin(client)
    assert "sos el admin" in r.get_data(as_text=True)
    # el setup ya no está disponible
    assert client.get("/setup").headers["Location"].endswith("/login")


def test_filas_y_columnas_iniciales(client):
    crear_admin(client)
    assert list(ids("filas")) == list(montos.FILAS_INICIALES)
    assert list(ids("columnas")) == ["Paco", "Antonio", "Ibra"]


def test_login(client):
    crear_admin(client)
    post(client, "/salir")
    assert client.get("/").headers["Location"].endswith("/login")
    assert "incorrectos" in ingresar(client, "admin", "mal").get_data(as_text=True)
    assert "Cajas por turno" in ingresar(client, "admin", "secreto1").get_data(as_text=True)


def test_post_sin_csrf_es_rechazado(client):
    crear_admin(client)
    assert client.post("/admin/filas", data={"nombre": "X"}).status_code == 400


def test_admin_agrega_y_quita_filas_y_columnas(client):
    crear_admin(client)
    post(client, "/admin/filas", nombre="Efectivo")
    post(client, "/admin/columnas", nombre="Retiros", signo="-1")
    assert "Efectivo" in ids("filas") and "Retiros" in ids("columnas")
    assert "Ya existe" in post(client, "/admin/filas", nombre="efectivo").get_data(as_text=True)

    # sin montos: se elimina de verdad
    post(client, f"/admin/filas/{ids('filas')['Efectivo']}/eliminar")
    assert "Efectivo" not in ids("filas")

    # con montos: se oculta y sus montos siguen contando en los saldos
    cargar(client, "2026-10-08", "noche", Mercado__Paco="1000", Mercado__Retiros="300")
    post(client, f"/admin/columnas/{ids('columnas')['Retiros']}/eliminar")
    with montos.app.app_context():
        r = montos.get_db().execute("SELECT activo FROM columnas WHERE nombre = 'Retiros'").fetchone()
    assert r["activo"] == 0
    d = caja("2026-10-08", "manana")
    assert "Retiros" not in [c["nombre"] for c in d["columnas"]]
    assert d["totales"]["anterior"] == 70000
    # en la caja donde tiene montos se sigue viendo
    assert "Retiros" in [c["nombre"] for c in caja("2026-10-08", "noche")["columnas"]]

    # volver a agregarla con el mismo nombre la reactiva
    post(client, "/admin/columnas", nombre="retiros")
    assert "Retiros" in [c["nombre"] for c in caja("2026-10-08", "manana")["columnas"]]


def test_usuario_ve_y_carga_solo_lo_asignado(client):
    crear_admin(client)
    post(client, "/admin/usuarios", nombre="paco", clave="clave123")
    post(client, "/admin/usuarios", nombre="ana", clave="clave456")
    f, c = ids("filas"), ids("columnas")
    uid = ids("usuarios")["paco"]
    post(client, f"/admin/usuarios/{uid}/permisos", filas=[f["Mercado"], f["Lemon"]], columnas=[c["Paco"]])

    ingresar(client, "paco", "clave123")
    pagina = client.get("/?fecha=2026-10-08&turno=noche").get_data(as_text=True)
    assert "Mercado" in pagina and "Lemon" in pagina and "Brubank" not in pagina
    assert "Antonio" not in pagina
    assert client.get("/admin").status_code == 403

    cargar(client, "2026-10-08", "noche", Mercado__Paco="1.500", Brubank__Paco="999", Mercado__Antonio="999")
    with montos.app.app_context():
        celdas = montos.get_db().execute("SELECT * FROM celdas").fetchall()
    assert [(r["fila_id"], r["columna_id"], r["centavos"], r["actualizado_por"]) for r in celdas] == [
        (f["Mercado"], c["Paco"], 150000, "paco")  # lo no asignado se ignora
    ]

    # un usuario sin asignaciones no ve nada
    ingresar(client, "ana", "clave456")
    assert "todavía no te asignó" in client.get("/").get_data(as_text=True)


def test_niveles_y_nombres(client):
    crear_admin(client)
    post(client, "/admin/usuarios", nombre="lucia", clave="clave123", rol="encargado", nombre_visible="Lucía Gómez")
    post(client, "/admin/usuarios", nombre="pepe", clave="clave123")
    with montos.app.app_context():
        roles = {r["nombre"]: (r["rol"], r["nombre_visible"]) for r in montos.get_db().execute("SELECT * FROM usuarios")}
    assert roles == {"admin": ("admin", None), "lucia": ("encargado", "Lucía Gómez"), "pepe": ("cajero", None)}
    assert "Nivel inválido" in post(client, "/admin/usuarios", nombre="x123", clave="clave123", rol="jefe").get_data(as_text=True)

    # el admin cambia nivel y nombre de otro, pero no su propio nivel
    post(client, f"/admin/usuarios/{ids('usuarios')['pepe']}/datos", rol="encargado", nombre_visible="Pepe")
    r = post(client, f"/admin/usuarios/{ids('usuarios')['admin']}/datos", rol="cajero")
    assert "No podés cambiar tu propio nivel" in r.get_data(as_text=True)
    with montos.app.app_context():
        db = montos.get_db()
        assert db.execute("SELECT rol, nombre_visible FROM usuarios WHERE nombre = 'pepe'").fetchone()[:] == ("encargado", "Pepe")
        assert db.execute("SELECT rol FROM usuarios WHERE nombre = 'admin'").fetchone()[0] == "admin"

    # cada uno elige su nombre en Mi cuenta, y es el que queda en los cambios
    ingresar(client, "lucia", "clave123")
    post(client, "/cuenta", nombre_visible="  Lu  ")
    assert "Lu · Encargado" in client.get("/cuenta").get_data(as_text=True)
    assert "La contraseña actual no es correcta" in post(
        client, "/cuenta", nombre_visible="Lu", clave_actual="mal", clave_nueva="nueva123"
    ).get_data(as_text=True)
    post(client, "/cuenta", nombre_visible="Lu", clave_actual="clave123", clave_nueva="nueva123")
    assert "Cajas por turno" in ingresar(client, "lucia", "nueva123").get_data(as_text=True)


def test_migra_es_admin_a_niveles(tmp_path, monkeypatch):
    ruta = tmp_path / "vieja.db"
    db = sqlite3.connect(ruta)
    db.executescript("""
        CREATE TABLE usuarios (id INTEGER PRIMARY KEY, nombre TEXT, clave_hash TEXT, es_admin INTEGER NOT NULL DEFAULT 0);
        INSERT INTO usuarios (nombre, clave_hash, es_admin) VALUES ('jefe', 'x', 1), ('caja', 'x', 0);
        PRAGMA user_version = 5;
    """)
    db.close()
    montos.app.config.update(DATABASE=str(ruta))
    with montos.app.app_context():
        r = montos.get_db().execute("SELECT nombre, rol FROM usuarios ORDER BY id").fetchall()
    assert [tuple(x) for x in r] == [("jefe", "admin"), ("caja", "cajero")]


def test_eliminar_usuario_y_no_al_admin(client):
    crear_admin(client)
    post(client, "/admin/usuarios", nombre="juan", clave="clave123")
    post(client, f"/admin/usuarios/{ids('usuarios')['juan']}/eliminar")
    assert "juan" not in ids("usuarios")
    assert post(client, f"/admin/usuarios/{ids('usuarios')['admin']}/eliminar").status_code == 404


# --- Cajas y cálculos ----------------------------------------------------------

def test_mover_entre_cajas():
    d = date(2026, 10, 8)
    assert montos.mover((d, 0), 1) == (d, 1)
    assert montos.mover((d, 2), 1) == (date(2026, 10, 9), 0)   # tarde -> noche del día siguiente
    assert montos.mover((d, 0), -1) == (date(2026, 10, 7), 2)  # noche -> tarde del día anterior
    assert montos.nombre_caja((d, 1)) == "Jueves 08/10/2026 – Mañana"


def test_turno_anterior_es_el_turno_actual_de_la_caja_previa(client):
    crear_admin(client)
    cargar(client, "2026-10-08", "noche", Mercado__Paco="150.000", Mercado__Antonio="80.000", Lemon__Ibra="-5.000")
    cargar(client, "2026-10-08", "manana", Mercado__Paco="45.000")
    noche, manana, tarde = (caja("2026-10-08", t) for t in ("noche", "manana", "tarde"))
    assert (noche["totales"]["anterior"], noche["totales"]["actual"]) == (0, 22500000)
    assert (manana["totales"]["anterior"], manana["totales"]["actual"]) == (22500000, 4500000)
    assert (tarde["totales"]["anterior"], tarde["totales"]["actual"]) == (4500000, 0)  # no acumula
    # si la caja anterior no tuvo montos, arranca en 0
    assert caja("2026-10-09", "manana")["totales"]["anterior"] == 0
    # corregir una caja cambia el turno anterior de la siguiente
    cargar(client, "2026-10-08", "manana", Mercado__Paco="50.000")
    assert caja("2026-10-08", "tarde")["totales"]["anterior"] == 5000000


def test_cuenta_de_la_caja(client):
    """turno anterior − (turno actual + bajada) = resultado; resultado + saldo = final; final − bono = total final."""
    crear_admin(client)
    cargar(client, "2026-10-08", "noche", Mercado__Paco="512.000")
    cargar(client, "2026-10-08", "manana", Mercado__Paco="57.000", Lemon__Antonio="13.500", Lemon__Ibra="-9.749,50")
    post(client, "/guardar", fecha="2026-10-08", turno="manana",
         x_bajada="150.000", x_deposito="200.000", x_retiro="35.000", x_bono="12.500")
    assert caja("2026-10-08", "manana")["totales"] == {
        "anterior": 51200000, "actual": 6075050, "bajada": 15000000,
        "resultado": 30124950, "saldo": 16500000, "final": 46624950, "bono": 1250000, "total_final": 45374950,
    }
    pagina = client.get("/?fecha=2026-10-08&turno=manana").get_data(as_text=True)
    assert "$301.249,50" in pagina and "$466.249,50" in pagina and "$453.749,50" in pagina
    assert pagina.index("Cuentas de la caja") < pagina.index("Otros datos del turno")

    with montos.app.app_context():
        h = montos.historial(montos.get_db())
    assert [(x["caja"][1], x["final"]) for x in h] == [(1, 46624950), (0, -51200000)]


def test_usuario_ve_los_totales_de_toda_la_caja(client):
    crear_admin(client)
    post(client, "/admin/usuarios", nombre="cajero", clave="clave123")
    f, c = ids("filas"), ids("columnas")
    post(client, f"/admin/usuarios/{ids('usuarios')['cajero']}/permisos", filas=[f["Lemon"]], columnas=[c["Ibra"]])
    cargar(client, "2026-10-08", "noche", Mercado__Paco="1000")
    cargar(client, "2026-10-08", "manana", Mercado__Antonio="300")
    ingresar(client, "cajero", "clave123")
    cargar(client, "2026-10-08", "manana", Lemon__Ibra="50")
    d = caja("2026-10-08", "manana", usuario="cajero")
    assert (d["totales"]["anterior"], d["totales"]["actual"]) == (100000, 35000)
    assert d["total_grilla"] == 5000
    pagina = client.get("/?fecha=2026-10-08&turno=manana").get_data(as_text=True)
    assert "son de toda la caja" in pagina and "Mercado" not in pagina


def test_columna_que_resta(client):
    crear_admin(client)
    post(client, "/admin/columnas", nombre="Retiros", signo="-1")
    cargar(client, "2026-10-08", "noche", Mercado__Paco="1000", Mercado__Retiros="250,50")
    d = caja("2026-10-08", "noche")
    assert d["total_columna"][ids("columnas")["Retiros"]] == 25050
    assert d["total_grilla"] == 74950
    assert d["totales"]["actual"] == 74950


def test_sin_parametros_muestra_el_turno_en_curso(client):
    crear_admin(client)
    pagina = client.get("/").get_data(as_text=True)
    assert "Jueves 08/10/2026 – Noche" in pagina and "Turno en curso" in pagina


@pytest.mark.parametrize("hora,caja_esperada", [
    ("2026-10-08 06:59", ("2026-10-08", 0)),  # noche: 23 a 7
    ("2026-10-08 07:00", ("2026-10-08", 1)),  # mañana: 7 a 15
    ("2026-10-08 14:59", ("2026-10-08", 1)),
    ("2026-10-08 15:00", ("2026-10-08", 2)),  # tarde: 15 a 23
    ("2026-10-08 22:59", ("2026-10-08", 2)),
    ("2026-10-08 23:00", ("2026-10-09", 0)),  # la noche ya es del día siguiente
])
def test_caja_en_curso(hora, caja_esperada):
    momento = datetime.strptime(hora, "%Y-%m-%d %H:%M").replace(tzinfo=montos.ZONA_HORARIA)
    assert montos.caja_en_curso(momento) == (date.fromisoformat(caja_esperada[0]), caja_esperada[1])


def test_dia_cerrado():
    def momento(texto):
        return datetime.strptime(texto, "%Y-%m-%d %H:%M").replace(tzinfo=montos.ZONA_HORARIA)
    jueves = date(2026, 10, 8)
    assert not montos.dia_cerrado((jueves, 0), momento("2026-10-08 22:59"))
    assert montos.dia_cerrado((jueves, 0), momento("2026-10-08 23:00"))
    assert montos.dia_cerrado((jueves, 2), momento("2026-10-09 10:00"))
    assert not montos.dia_cerrado((date(2026, 10, 9), 0), momento("2026-10-08 23:30"))  # noche del viernes en curso


def test_dia_terminado_solo_lo_cambia_el_admin(client):
    crear_admin(client)
    post(client, "/admin/usuarios", nombre="cajero", clave="clave123", rol="cajero")
    f, c = ids("filas"), ids("columnas")
    post(client, f"/admin/usuarios/{ids('usuarios')['cajero']}/permisos", filas=[f["Mercado"]], columnas=[c["Paco"]])

    ingresar(client, "cajero", "clave123")
    r = cargar(client, "2026-10-07", "tarde", Mercado__Paco="100")  # día anterior: cerrado
    assert "solo el admin puede hacer cambios" in r.get_data(as_text=True)
    assert caja("2026-10-07", "tarde")["totales"]["actual"] == 0
    pagina = client.get("/?fecha=2026-10-07&turno=tarde").get_data(as_text=True)
    assert "readonly" in pagina and "Este día ya terminó" in pagina
    cargar(client, "2026-10-08", "manana", Mercado__Paco="100")  # el día de hoy sigue abierto
    assert caja("2026-10-08", "manana")["totales"]["actual"] == 10000

    ingresar(client, "admin", "secreto1")
    cargar(client, "2026-10-07", "tarde", Mercado__Paco="100")
    assert caja("2026-10-07", "tarde")["totales"]["actual"] == 10000


def test_vaciar_celda_la_borra(client):
    crear_admin(client)
    cargar(client, "2026-10-08", "noche", Mercado__Paco="10", Lemon__Paco="20")
    cargar(client, "2026-10-08", "noche", Lemon__Paco="")
    with montos.app.app_context():
        assert montos.get_db().execute("SELECT COUNT(*) FROM celdas").fetchone()[0] == 1


def test_error_de_monto_se_informa(client):
    crear_admin(client)
    r = cargar(client, "2026-10-08", "noche", Prex__Ibra="abc")
    assert "Prex / Ibra" in r.get_data(as_text=True)


def test_guardar_sin_caja_valida(client):
    crear_admin(client)
    assert post(client, "/guardar", fecha="mal", turno="noche").status_code == 400


def test_otros_datos_saldo_y_total_con_bajada(client):
    crear_admin(client)
    cargar(client, "2026-10-08", "noche", Mercado__Paco="1000")
    post(client, "/guardar", fecha="2026-10-08", turno="manana",
         x_deposito="2.000", x_retiro="700", x_bajada="500", x_bono="1.200,50", x_saldo="999")
    noche, manana = caja("2026-10-08", "noche"), caja("2026-10-08", "manana")
    # el saldo no se guarda aunque lo manden: se calcula
    assert {k: v["centavos"] for k, v in manana["extras"].items()} == {
        "deposito": 200000, "retiro": 70000, "bajada": 50000, "bono": 120050
    }
    # no suman en el turno actual; se usan en el resultado y el final
    t = manana["totales"]
    assert (t["actual"], t["bajada"], t["saldo"]) == (0, 50000, 130000)
    assert caja("2026-10-08", "tarde")["totales"]["anterior"] == 0
    assert noche["extras"] == {} and noche["totales"]["saldo"] == 0  # son de cada caja

    pagina = client.get("/?fecha=2026-10-08&turno=manana").get_data(as_text=True)
    assert 'value="1.200,50"' in pagina and "$1.300,00" in pagina
    assert 'name="x_saldo"' not in pagina
    orden = [pagina.index(f'name="x_{campo}"') for campo in ("deposito", "retiro", "bono", "bajada")]
    assert orden == sorted(orden)

    # vaciar un campo lo borra; un valor inválido se informa
    r = post(client, "/guardar", fecha="2026-10-08", turno="manana", x_retiro="", x_bono="abc")
    assert "Bono: «abc»" in r.get_data(as_text=True)
    manana = caja("2026-10-08", "manana")
    assert set(manana["extras"]) == {"deposito", "bono", "bajada"} and manana["totales"]["saldo"] == 200000


def test_exportar_excel(client):
    crear_admin(client)
    cargar(client, "2026-10-08", "noche", Mercado__Paco="1000")
    cargar(client, "2026-10-08", "manana", Mercado__Antonio="250")
    r = client.get("/exportar.xlsx?fecha=2026-10-08&turno=manana")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb["Caja"]
    assert ws["A1"].value == "Caja: Jueves 08/10/2026 – Mañana"
    assert [ws.cell(row=3, column=i).value for i in range(1, 6)] == ["Cuenta", "Paco", "Antonio", "Ibra", None]
    assert ws["A4"].value == "Mercado" and ws["C4"].value == 250 and ws["E4"].value is None
    assert ws["A13"].value == "Total" and ws["D13"].value == "=SUM(D4:D12)" and ws["E13"].value is None
    assert [(ws[f"A{i}"].value, ws[f"B{i}"].value) for i in range(15, 23)] == [
        ("Turno anterior", 1000), ("Turno actual", "=+B13+C13+D13"), ("Bajada", "=B29"),
        ("Resultado", "=B15-(B16+B17)"), ("Saldo", "=B28"), ("Final", "=B18+B19"),
        ("Bono", "=B27"), ("Total final", "=B20-B21"),
    ]
    assert ws["A24"].value == "Otros datos del turno"
    assert [(ws[f"A{i}"].value, ws[f"B{i}"].value) for i in range(25, 30)] == [
        ("Depósito", None), ("Retiro", None), ("Bono", None), ("Saldo", "=B25-B26"), ("Bajada", None)
    ]
    hs = wb["Historial"]
    assert [c.value for c in hs[3]][1:] == ["Mañana", 1000, 250, 0, 750, 0, 750, 0, 750]

    post(client, "/guardar", fecha="2026-10-08", turno="manana", x_bono="75")
    ws = load_workbook(io.BytesIO(client.get("/exportar.xlsx?fecha=2026-10-08&turno=manana").data))["Caja"]
    assert ws["B27"].value == 75 and ws["B22"].value == "=B20-B21"
